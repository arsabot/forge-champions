import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen slides (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    c_orange = RGBColor(0xFF, 0xA4, 0x00)
    c_deep_orange = RGBColor(0xFF, 0x6B, 0x00)
    c_dark = RGBColor(0x22, 0x1F, 0x20)
    c_beige = RGBColor(0xF8, 0xF5, 0xEE)
    c_sand = RGBColor(0xEE, 0xE9, 0xDF)
    c_white = RGBColor(0xFF, 0xFF, 0xFF)
    c_gray = RGBColor(0x6E, 0x6D, 0x7A)
    c_teal = RGBColor(0x00, 0xA8, 0x96)

    def add_header(slide, title_text, category_text="FUNDACIÓN FORGE • TEAM 7"):
        # Header banner shape
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_deep_orange
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = c_dark
        p_title.space_before = Pt(2)

    # ----------------------------------------------------
    # SLIDE 1: PORTADA & EQUIPO TEAM 7
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    
    # Background fill
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_beige
    bg1.line.color.rgb = c_beige

    # Top accent bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = c_deep_orange
    top_bar.line.fill.background()

    # Left Column: Title & Subtitle
    left_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(6.2), Inches(4.5))
    tf1 = left_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "INICIATIVA OFICIAL"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = c_deep_orange
    
    p2 = tf1.add_paragraph()
    p2.text = "Forge Champions"
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = c_dark
    p2.space_before = Pt(8)
    
    p3 = tf1.add_paragraph()
    p3.text = "Portal MVP de Masterclasses, Mentoría e Impacto Social para Jóvenes de Latinoamérica."
    p3.font.size = Pt(15)
    p3.font.color.rgb = c_gray
    p3.space_before = Pt(12)

    p_badge = tf1.add_paragraph()
    p_badge.text = "• Despliegue en AWS EC2 • Next.js 15 • Supabase PostgreSQL • Google Calendar"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = c_dark
    p_badge.space_before = Pt(20)

    # Right Card: Team 7 Members (Alphabetical)
    team_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.3), Inches(4.8), Inches(5.0))
    team_card.fill.solid()
    team_card.fill.fore_color.rgb = c_white
    team_card.line.color.rgb = c_sand
    team_card.line.width = Pt(1.5)

    team_box = s1.shapes.add_textbox(Inches(7.8), Inches(1.6), Inches(4.2), Inches(4.4))
    tft = team_box.text_frame
    tft.word_wrap = True

    pt_title = tft.paragraphs[0]
    pt_title.text = "PARTICIPANTES — TEAM 7"
    pt_title.font.size = Pt(14)
    pt_title.font.bold = True
    pt_title.font.color.rgb = c_deep_orange

    pt_sub = tft.add_paragraph()
    pt_sub.text = "Integrantes en Orden Alfabético:"
    pt_sub.font.size = Pt(11)
    pt_sub.font.color.rgb = c_gray
    pt_sub.space_before = Pt(4)

    integrantes = [
        "Juan Gimenez",
        "Lautaro Pavan",
        "Mijail Arispe",
        "Rodrigo Saavedra",
        "Roy Fiorilo"
    ]

    for i, name in enumerate(integrantes, 1):
        pi = tft.add_paragraph()
        pi.text = f"{i}. {name}"
        pi.font.size = Pt(15)
        pi.font.bold = True
        pi.font.color.rgb = c_dark
        pi.space_before = Pt(10)

    # ----------------------------------------------------
    # SLIDE 2: HERO SECTION & PROPUESTA DE VALOR
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Hero Section & Llamados a la Acción", "EXPERIENCIA DE USUARIO")
    
    # Text Description Box (Left)
    t2_box = s2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf2 = t2_box.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "Estructura de Alta Conversión"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_dark
    
    p = tf2.add_paragraph()
    p.text = "• Identidad Oficial de Fundación Forge: Paleta cálida (#FFA400, #FF6B00, #221F20, #F8F5EE) y tipografía moderna.\n\n• 3 Preguntas Clave ('¿Querés...?') que interpelan directamente a profesionales y líderes corporativos.\n\n• Mensaje de Cierre Motivador: '¡Estás en el lugar indicado! ¡Podés ser un Champion!'.\n\n• Doble CTA Segmentado: 'Postular mi Empresa' (RSC / Alianza) y 'Quiero ser Champion' (Speakers)."
    p.font.size = Pt(12)
    p.font.color.rgb = c_gray
    p.space_before = Pt(10)

    # Image Screenshot (Right)
    if os.path.exists('screenshots/02_hero_section.png'):
        s2.shapes.add_picture('screenshots/02_hero_section.png', Inches(5.6), Inches(1.5), width=Inches(6.9))

    # ----------------------------------------------------
    # SLIDE 3: CREACIÓN DE VALOR (PILAR CORPORATIVO & PROFESIONAL)
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Creación de Valor (Impacto de Doble Vía)", "PROPUESTA DE VALOR")
    
    t3_box = s3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf3 = t3_box.text_frame
    tf3.word_wrap = True
    
    p = tf3.paragraphs[0]
    p.text = "Articulación Corporativa e Institucional"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_dark
    
    p = tf3.add_paragraph()
    p.text = "• Para la Empresa:\n  - Difusión de Marca Empleadora (Employer Branding).\n  - Impacto RSC a Costo Cero con métricas sociales reales.\n  - Formación en oratoria y liderazgo para equipos internos.\n\n• Para el Profesional (Champion):\n  - Posicionamiento como referente del sector tech / corporativo.\n  - Networking con propósito y conexión con el futuro talento.\n  - Certificación e insignia oficial emitida por Fundación Forge."
    p.font.size = Pt(12)
    p.font.color.rgb = c_gray
    p.space_before = Pt(10)

    if os.path.exists('screenshots/03_propuesta_valor.png'):
        s3.shapes.add_picture('screenshots/03_propuesta_valor.png', Inches(5.6), Inches(1.5), width=Inches(6.9))

    # ----------------------------------------------------
    # SLIDE 4: CHAMPIONS SESSIONS (CARRUSEL INTERACTIVO)
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Champions Sessions (Muro de Masterclasses)", "CONTENIDO & REPRODUCCIÓN")
    
    t4_box = s4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf4 = t4_box.text_frame
    tf4.word_wrap = True
    
    p = tf4.paragraphs[0]
    p.text = "Archivo de Charlas y Snippets de Video"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_dark
    
    p = tf4.add_paragraph()
    p.text = "• Primera Masterclass Destacada: Mariana Albarracín (VP de Ciberseguridad & Risk en J.P. Morgan).\n\n• Charla TED de Ejemplo: Susan Cain (El poder de la autenticidad y el liderazgo).\n\n• Líderes de Mercado Libre, Google Cloud, Globant y Nubank.\n\n• Carrusel Interactivo por Índice: Botones de avance/retroceso, contador '1 / 6' y puntos indicadores de salto directo."
    p.font.size = Pt(12)
    p.font.color.rgb = c_gray
    p.space_before = Pt(10)

    if os.path.exists('screenshots/04_champions_sessions.png'):
        s4.shapes.add_picture('screenshots/04_champions_sessions.png', Inches(5.6), Inches(1.6), width=Inches(6.9))

    # ----------------------------------------------------
    # SLIDE 5: CALENDARIO OPERATIVO & GOOGLE CALENDAR
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Calendario Operativo & Google Calendar", "INTEGRACIÓN EN TIEMPO REAL")
    
    t5_box = s5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf5 = t5_box.text_frame
    tf5.word_wrap = True
    
    p = tf5.paragraphs[0]
    p.text = "Sincronización y Acceso en Vivo"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_dark
    
    p = tf5.add_paragraph()
    p.text = "• Botón '+ Google Calendar' en 1 Clic:\n  Genera el evento automáticamente con título, fecha, speaker, cargo y link de Zoom.\n\n• Seguridad y Desbloqueo Inteligente de Zoom:\n  El botón de sala en vivo se habilita únicamente 5 minutos antes de la hora pactada para evitar ingresos indebidos.\n\n• Consultas dinámicas a Supabase para masterclasses programadas."
    p.font.size = Pt(12)
    p.font.color.rgb = c_gray
    p.space_before = Pt(10)

    if os.path.exists('screenshots/05_calendario_operativo.png'):
        s5.shapes.add_picture('screenshots/05_calendario_operativo.png', Inches(5.6), Inches(1.6), width=Inches(6.9))

    # ----------------------------------------------------
    # SLIDE 6: FORMULARIO DE CONVOCATORIA & SUPABASE
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. Formulario de Postulación & Base de Datos", "CAPTACIÓN & PERSISTENCIA")
    
    t6_box = s6.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf6 = t6_box.text_frame
    tf6.word_wrap = True
    
    p = tf6.paragraphs[0]
    p.text = "Segmentación de Alta Conversión"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_dark
    
    p = tf6.add_paragraph()
    p.text = "• Tabs de Segmentación:\n  - 'Soy Profesional (Champion)'\n  - 'Represento a una Empresa'\n\n• Validaciones en Tiempo Real:\n  Captura de nombre, empresa actual, cargo, LinkedIn y propuesta temática.\n\n• Base de Datos Relacional PostgreSQL (Supabase):\n  Tablas 'postulaciones_champion' y 'sesiones_masterclass' con ENUMs y políticas de seguridad RLS."
    p.font.size = Pt(12)
    p.font.color.rgb = c_gray
    p.space_before = Pt(10)

    if os.path.exists('screenshots/06_formulario_postulacion.png'):
        s6.shapes.add_picture('screenshots/06_formulario_postulacion.png', Inches(5.6), Inches(1.6), width=Inches(6.9))

    # ----------------------------------------------------
    # SLIDE 7: CIERRE, EQUIPO & ENLACES
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    
    bg7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = c_beige
    bg7.line.color.rgb = c_beige

    top_bar7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    top_bar7.fill.solid()
    top_bar7.fill.fore_color.rgb = c_deep_orange
    top_bar7.line.fill.background()

    cbox = s7.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    tfc = cbox.text_frame
    tfc.word_wrap = True

    p = tfc.paragraphs[0]
    p.text = "¡Muchas Gracias!"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_dark

    p = tfc.add_paragraph()
    p.text = "Presentación de Proyecto — Forge Champions (Fundación Forge)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_deep_orange
    p.space_before = Pt(6)

    p = tfc.add_paragraph()
    p.text = "Participantes Team 7 (Orden Alfabético):"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = c_dark
    p.space_before = Pt(20)

    for name in integrantes:
        pi = tfc.add_paragraph()
        pi.text = f"• {name}"
        pi.font.size = Pt(14)
        pi.font.color.rgb = c_gray
        pi.space_before = Pt(4)

    p_links = tfc.add_paragraph()
    p_links.text = "🌐 Servidor en Vivo: http://ec2-52-90-35-158.compute-1.amazonaws.com:3000\n📂 Repositorio GitHub: https://github.com/cfgba26/Team-7"
    p_links.font.size = Pt(13)
    p_links.font.bold = True
    p_links.font.color.rgb = c_deep_orange
    p_links.space_before = Pt(24)

    output_path = 'Forge_Champions_Team7_Presentacion.pptx'
    prs.save(output_path)
    print(f"Presentation saved successfully as {output_path}")

if __name__ == '__main__':
    build_presentation()
